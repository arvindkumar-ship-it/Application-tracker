from pathlib import Path
import pymupdf as fitz
from jinja2 import Environment,FileSystemLoader,select_autoescape
from docx import Document
from docx.shared import Inches,Pt
from playwright.async_api import async_playwright
from app.config import get_settings


def make_docx(path,sections,contacts):
    document=Document();style=document.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10.5)
    for section in document.sections:
        section.top_margin=Inches(.55);section.bottom_margin=Inches(.55);section.left_margin=Inches(.63);section.right_margin=Inches(.63)
    document.add_heading(sections['profile']['full_name'],0);document.add_paragraph(' · '.join(contacts))
    for name in sections['section_order']:
        value=sections.get(name)
        if not value: continue
        document.add_heading(name.title(),1)
        if name=='summary': document.add_paragraph(value['text'])
        elif name=='skills': document.add_paragraph(', '.join(x['name'] for x in value))
        else:
            for entry in value:
                paragraph=document.add_paragraph();paragraph.add_run(entry['title']).bold=True
                if entry.get('dates'): paragraph.add_run(' | '+entry['dates'])
                if entry.get('subtitle'): document.add_paragraph(entry['subtitle'])
                for bullet in entry['bullets']: document.add_paragraph(bullet['text'],style='List Bullet')
    document.save(path)


async def render(job_id,version,sections,max_pages):
    directory=Path(get_settings().data_dir)/'resumes'/str(job_id);directory.mkdir(parents=True,exist_ok=True)
    pdf_path=directory/f'v{version}.pdf';docx_path=directory/f'v{version}.docx'
    templates=Path(__file__).resolve().parent.parent/'templates'
    environment=Environment(loader=FileSystemLoader(templates),autoescape=select_autoescape(['html']))
    template=environment.get_template('resume_ats.html');css=(templates/'resume_ats.css').read_text()
    profile=sections['profile'];contacts=[str(profile[k]) for k in ['email','phone','location_city','linkedin_url','github_url','portfolio_url'] if profile.get(k)]
    warnings=[]
    async with async_playwright() as pw:
        import os
        # Containers use bundled Chromium. Windows development can use installed Chrome.
        local_channel='chrome' if os.name=='nt' and not Path(pw.chromium.executable_path).is_file() else None
        browser=await pw.chromium.launch(headless=True,args=['--no-sandbox'],channel=local_channel)
        page=await browser.new_page();await page.route('**/*',lambda route:route.abort())
        try:
            for _ in range(100):
                html=template.render(profile=profile,contacts=contacts,sections=sections,css=css)
                await page.set_content(html,wait_until='load');await page.pdf(path=str(pdf_path),format='A4',print_background=False)
                with fitz.open(pdf_path) as pdf: count=pdf.page_count
                if count<=max_pages: break
                candidates=[(bullet.get('score',0),name,i,j) for name in sections['section_order'] if name not in {'summary','skills'} for i,entry in enumerate(sections.get(name,[])) for j,bullet in enumerate(entry['bullets'])]
                if not candidates: raise ValueError('Resume exceeds the page limit with fixed profile/education content. Shorten the recorded content or allow two pages.')
                _,name,i,j=min(candidates);sections[name][i]['bullets'].pop(j);warnings.append('A low-priority bullet was removed to fit the page limit.')
            else: raise ValueError('Could not fit the resume within the page limit')
        finally: await browser.close()
    make_docx(docx_path,sections,contacts)
    return str(pdf_path),str(docx_path),count,warnings
